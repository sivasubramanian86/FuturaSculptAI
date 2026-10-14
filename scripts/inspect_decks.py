import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')

def inspect_presentation(path, max_slides=60):
    prs = pptx.Presentation(path)
    print(f"=== {path} ===")
    print(f"Slide Dimensions: {prs.slide_width.inches} x {prs.slide_height.inches} inches")
    print(f"Total Slides: {len(prs.slides)}")
    
    for idx, slide in enumerate(prs.slides):
        if idx >= max_slides:
            break
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t and t not in texts:
                        texts.append(t)
        header = texts[0] if texts else "BLANK"
        print(f"Slide {idx + 1:02d}: {header[:80]}")
        for t in texts[1:5]:
            print(f"   - {t[:100]}")

if __name__ == "__main__":
    print("INSPECTING TEMPLATE:")
    inspect_presentation(r"C:\Users\USER\Downloads\Submission Template _ AI Builder Cup.pptx", 16)
    print("\nINSPECTING REFERENCE PRESENTATION:")
    inspect_presentation(r"D:\Siva\Books\CAREER\HACKATHON\AWS_AI_For_Bharat_Harve_Logix_AI_H2S_Submissions\AWS_AI_for_Bharat_Hackathon_Prototype_Development_Submission_Team_Cybernauts_March_2026.pptx", 60)
